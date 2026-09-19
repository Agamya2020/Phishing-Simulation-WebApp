"""Users, Groups, Templates, Departments CRUD routers."""
import csv
import io

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import EmailStr, TypeAdapter, ValidationError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import delete, select, update
from app.core.db import get_db
from app.core.security import require_admin
from app.models.models import (
    Campaign,
    CampaignEvent,
    Department,
    Group,
    Template,
    User,
)
from app.schemas.schemas import (
    UserOut, UserCreate, UserUpdate,
    GroupOut, GroupCreate,
    TemplateOut, TemplateCreate, TemplateUpdate,
    DepartmentOut, DepartmentCreate,
)


email_validator = TypeAdapter(EmailStr)


def normalize_key(value):
    return str(value or "").strip().lower().replace("_", " ")


def value_from(row, *possible_names):
    normalized = {
        normalize_key(key): str(value or "").strip()
        for key, value in row.items()
    }

    for name in possible_names:
        value = normalized.get(normalize_key(name))
        if value:
            return value

    return ""


def parse_user_file(filename: str, content: bytes):
    filename = filename.lower()

    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return list(reader)

    if filename.endswith(".xlsx"):
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))

        if not rows:
            return []

        headers = [str(value or "").strip() for value in rows[0]]
        return [
            dict(zip(headers, row))
            for row in rows[1:]
            if any(value is not None for value in row)
        ]

    if filename.endswith(".txt"):
        text = content.decode("utf-8-sig")
        records = []

        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue

            if "\t" in line:
                parts = [part.strip() for part in line.split("\t")]
            elif "," in line:
                parts = [part.strip() for part in line.split(",")]
            else:
                parts = [line]

            if len(parts) >= 2:
                records.append({"name": parts[0], "email": parts[1]})
            else:
                records.append({"name": "", "email": parts[0]})

        return records

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Use CSV, XLSX or TXT.",
    )

# ─── Users ─────────────────────────────────────────────────────────────────────
users_router = APIRouter(
    prefix="/users",
    tags=["users"],
    dependencies=[
        Depends(require_admin)
    ],
)


@users_router.get("", response_model=list[UserOut])
async def list_users(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(User).order_by(User.name))
    return result.scalars().all()


@users_router.post("/import")
async def import_users(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    filename = file.filename or ""

    if not filename.lower().endswith((".csv", ".xlsx", ".txt")):
        raise HTTPException(
            status_code=400,
            detail="Upload a CSV, XLSX or TXT file.",
        )

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail="File must be smaller than 5 MB.",
        )

    rows = parse_user_file(filename, content)

    department_result = await db.execute(select(Department))
    departments = department_result.scalars().all()
    department_lookup = {}

    for department in departments:
        department_lookup[department.id.lower()] = department.id
        department_lookup[department.name.lower()] = department.id
        department_lookup[department.code.lower()] = department.id

    existing_result = await db.execute(select(User.email))
    existing_emails = {
        str(email).lower()
        for email in existing_result.scalars()
    }

    imported = 0
    skipped = 0
    errors = []

    for row_number, row in enumerate(rows, start=2):
        email = value_from(
            row,
            "email",
            "email address",
            "mail",
            "e-mail",
        ).lower()
        name = value_from(
            row,
            "name",
            "full name",
            "employee name",
        )
        role = value_from(row, "role", "designation") or "employee"
        status = value_from(row, "status").lower() or "active"
        department_value = value_from(
            row,
            "department",
            "department name",
            "department code",
            "department id",
        )

        if not email:
            skipped += 1
            errors.append(f"Row {row_number}: missing email.")
            continue

        try:
            email = str(email_validator.validate_python(email)).lower()
        except ValidationError:
            skipped += 1
            errors.append(f"Row {row_number}: invalid email.")
            continue

        if email in existing_emails:
            skipped += 1
            continue

        if not name:
            name = (
                email.split("@", 1)[0]
                .replace(".", " ")
                .replace("_", " ")
                .title()
            )

        department_id = None

        if department_value:
            department_id = department_lookup.get(department_value.lower())

            if not department_id:
                errors.append(
                    f"Row {row_number}: department '{department_value}' "
                    "was not found; user imported without department."
                )

        user = User(
            name=name,
            email=email,
            role=role,
            status=status,
            department_id=department_id,
        )
        db.add(user)
        existing_emails.add(email)
        imported += 1

    await db.commit()

    return {
        "ok": True,
        "imported": imported,
        "skipped": skipped,
        "errors": errors,
    }


@users_router.get("/{user_id}", response_model=UserOut)
async def get_user(user_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(User).where(User.id == user_id))
    user = result.scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user


@users_router.post("", response_model=UserOut)
async def create_user(payload: UserCreate, db: AsyncSession = Depends(get_db)):
    user = User(**payload.model_dump())
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


@users_router.patch("/{user_id}", response_model=UserOut)
async def update_user(user_id: str, payload: UserUpdate, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(User).where(User.id == user_id))
    user = result.scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    update_data = payload.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(user, field, value)
    await db.commit()
    await db.refresh(user)
    return user


@users_router.delete("/{user_id}")
async def delete_user(user_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(User).where(User.id == user_id))
    user = result.scalar_one_or_none()

    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    result = await db.execute(select(Group))
    groups = result.scalars().all()

    for group in groups:
        if user_id in (group.member_ids or []):
            group.member_ids = [
                item for item in group.member_ids if item != user_id
            ]

    result = await db.execute(select(Campaign))
    campaigns = result.scalars().all()
    affected_campaigns = []

    for campaign in campaigns:
        if user_id in (campaign.target_user_ids or []):
            campaign.target_user_ids = [
                item
                for item in campaign.target_user_ids
                if item != user_id
            ]
            campaign.target_count = len(campaign.target_user_ids)
            affected_campaigns.append(campaign)

    await db.execute(
        delete(CampaignEvent).where(CampaignEvent.user_id == user_id)
    )
    await db.flush()

    for campaign in affected_campaigns:
        event_result = await db.execute(
            select(CampaignEvent).where(
                CampaignEvent.campaign_id == campaign.id
            )
        )
        events = event_result.scalars().all()

        def count_event(event_type):
            return len({
                event.user_id or event.user_email
                for event in events
                if event.event_type == event_type
            })

        campaign.open_count = count_event("opened")
        campaign.click_count = count_event("clicked")
        campaign.report_count = count_event("reported")
        campaign.creds_count = count_event("creds_entered")

    await db.delete(user)
    await db.commit()
    return {"ok": True, "message": "User deleted."}


# ─── Groups ────────────────────────────────────────────────────────────────────
groups_router = APIRouter(
    prefix="/groups",
    tags=["groups"],
    dependencies=[
        Depends(require_admin)
    ],
)


@groups_router.get("", response_model=list[GroupOut])
async def list_groups(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Group).order_by(Group.name))
    return result.scalars().all()


@groups_router.post("", response_model=GroupOut)
async def create_group(payload: GroupCreate, db: AsyncSession = Depends(get_db)):
    group = Group(**payload.model_dump())
    db.add(group)
    await db.commit()
    await db.refresh(group)
    return group


@groups_router.patch("/{group_id}", response_model=GroupOut)
async def update_group(group_id: str, payload: GroupCreate, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Group).where(Group.id == group_id))
    group = result.scalar_one_or_none()
    if not group:
        raise HTTPException(status_code=404, detail="Group not found")
    for field, value in payload.model_dump(exclude_none=True).items():
        setattr(group, field, value)
    await db.commit()
    await db.refresh(group)
    return group


@groups_router.delete("/{group_id}")
async def delete_group(group_id: str, db: AsyncSession = Depends(get_db)):
    await db.execute(delete(Group).where(Group.id == group_id))
    await db.commit()
    return {"ok": True}


# ─── Templates ─────────────────────────────────────────────────────────────────
templates_router = APIRouter(
    prefix="/templates",
    tags=["templates"],
    dependencies=[
        Depends(require_admin)
    ],
)


@templates_router.get("", response_model=list[TemplateOut])
async def list_templates(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Template).order_by(Template.created_at.desc()))
    return result.scalars().all()


@templates_router.get("/{template_id}", response_model=TemplateOut)
async def get_template(template_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Template).where(Template.id == template_id))
    template = result.scalar_one_or_none()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    return template


@templates_router.post("", response_model=TemplateOut)
async def create_template(payload: TemplateCreate, db: AsyncSession = Depends(get_db)):
    template = Template(**payload.model_dump())
    db.add(template)
    await db.commit()
    await db.refresh(template)
    return template


@templates_router.patch("/{template_id}", response_model=TemplateOut)
async def update_template(template_id: str, payload: TemplateUpdate, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Template).where(Template.id == template_id))
    template = result.scalar_one_or_none()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    update_data = payload.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(template, field, value)
    await db.commit()
    await db.refresh(template)
    return template


@templates_router.delete("/{template_id}")
async def delete_template(template_id: str, db: AsyncSession = Depends(get_db)):
    await db.execute(delete(Template).where(Template.id == template_id))
    await db.commit()
    return {"ok": True}


# ─── Departments ───────────────────────────────────────────────────────────────
departments_router = APIRouter(
    prefix="/departments",
    tags=["departments"],
    dependencies=[
        Depends(require_admin)
    ],
)


@departments_router.get("", response_model=list[DepartmentOut])
async def list_departments(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Department).order_by(Department.name))
    return result.scalars().all()


@departments_router.post("", response_model=DepartmentOut)
async def create_department(payload: DepartmentCreate, db: AsyncSession = Depends(get_db)):
    dept = Department(**payload.model_dump())
    db.add(dept)
    await db.commit()
    await db.refresh(dept)
    return dept


@departments_router.delete("/{dept_id}")
async def delete_department(dept_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(
        select(Department).where(Department.id == dept_id)
    )
    department = result.scalar_one_or_none()

    if not department:
        raise HTTPException(status_code=404, detail="Department not found")

    await db.execute(
        update(User)
        .where(User.department_id == dept_id)
        .values(department_id=None)
    )
    await db.execute(
        update(Group)
        .where(Group.department_id == dept_id)
        .values(department_id=None)
    )
    await db.delete(department)
    await db.commit()
    return {"ok": True, "message": "Department deleted."}
